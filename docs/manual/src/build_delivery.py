# -*- coding: utf-8 -*-
# Сборка книги ручного тестирования флоу доставки. Структура и оформление — как в build.py.
# Запуск: python3 build_delivery.py <папка с cases_delivery.py> <выходной .xlsx>
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.workbook.properties import CalcProperties

sys.path.insert(0, sys.argv[1])
from cases_delivery import CASES, BUGS, QUESTIONS

OUT = sys.argv[2]
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="3A3A3A")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
SUB = Font(name=FONT, size=11, bold=True)
NOTE = Font(name=FONT, size=9, italic=True, color="666666")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
SECTION_FILL = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)
GREEN, RED, ORANGE, GREY, BLUE, PURPLE = "C6EFCE", "FFC7CE", "FFE0B2", "E7E6E6", "DDEBF7", "E4DFEC"


def fill(hex_):
    return PatternFill("solid", fgColor=hex_)


def header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.alignment, c.border = HDR_FONT, HDR_FILL, CENTER, BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w


def body_row(ws, row, values, center_cols=()):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.border = BODY, BORDER
        c.alignment = CENTER if i in center_cols else WRAP


def key_value_sheet(ws, start_row, rows):
    r = start_row
    for k, v in rows:
        a = ws.cell(row=r, column=1, value=k)
        if v is None:
            a.font = SUB
            for col in (1, 2):
                ws.cell(row=r, column=col).fill = SECTION_FILL
        else:
            a.font = BOLD
            b = ws.cell(row=r, column=2, value=v)
            b.font, b.alignment = BODY, WRAP
        a.alignment = WRAP
        r += 1


ids = [c[0] for c in CASES]
assert len(ids) == len(set(ids)), "дубли ID кейсов"

wb = Workbook()

# ============================ 1. Как пользоваться ============================
ws = wb.active
ws.title = "Как пользоваться"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 110
ws["A1"] = "Tez Chakana — ручное тестирование флоу доставки: как пользоваться этим файлом"
ws["A1"].font = TITLE
key_value_sheet(ws, 3, [
    ("Листы", None),
    ("Тест-план", "Что проверяем, кто за какой ролью, соответствие экранов макета и статусов панели, порядок прогона по заказам, деньги и безопасность."),
    ("Тест-кейсы", "Все кейсы флоу доставки. Фильтры: «Модуль», «Набор», «Приоритет», «Риск». В шагах роль указана в скобках: [Клиент], [Мерчант], [Курьер], [Оператор]."),
    ("Прогон", "Шаблон одного прогона. На каждую сборку — копия листа (ПКМ по вкладке → «Переместить или скопировать» → «Создать копию»). Названия подтягиваются из «Тест-кейсы» по ID."),
    ("Баги", "Баги флоу доставки. Уже внесены 12 находок прогона 29.09 (заказ TEZ00871) — большинство со статусом «Перепроверить»."),
    ("Вопросы", "Вопросы к продукту и дизайну: ответ превращает вопрос в баг или в уточнение ожидаемого результата кейса."),
    ("Как проходить", None),
    ("По заказам", "Кейсы выполняются не по одному, а цепочкой на одном заказе — см. «Тест-план → Порядок прогона». Номер заказа пишите в «Комментарий» каждой строки прогона."),
    ("Роли", "Вы одновременно мерчант (панель под суперадмином, отдельный браузер или профиль) и курьер (телефон). Клиент — эмулятор или телефон. Оператор — панель под ролью оператора в Chrome."),
    ("Что записывать", "Если ожидаемый результат говорит «записать» — поведение в макете не описано: впишите факт в «Комментарий» и поставьте Pass, если поведение понятное и корректное."),
    ("Время обновления", "На переходах засекайте, через сколько секунд статус появился у клиента (RSL-06): ≤10 с — норма, 10–30 с — замечание, >30 с — баг."),
    ("Наборы (колонка «Набор»)", None),
    ("Smoke", "Основной путь на одном заказе с наличными: оформление → 01 → 02 → 03 → 04 → 05 → 06 и данные у оператора и курьера. ~30 мин."),
    ("Sanity", "Smoke + отклонение, замена, задержки, недоступный клиент, устойчивость — ключевые проверки каждого сценария."),
    ("Regression", "Все кейсы, включая редкие ветки и мелкие проверки макета."),
    ("Статусы в «Прогон»", None),
    ("Pass", "Факт совпал с ожидаемым."),
    ("Fail", "Не совпал → баг в листе «Баги», его ID — в колонку «Баг»."),
    ("Blocked", "Нельзя выполнить (неизвестен таймер, нет функции, упал предыдущий шаг) — причину в «Комментарий»."),
    ("Skip", "Сознательно пропущен."),
    ("Риск (колонка)", None),
    ("Безопасно", "Только просмотр."),
    ("Меняет данные", "Меняет тестовый заказ вручную (например, ручная смена статуса оператором)."),
    ("Реальный заказ", "Создаёт настоящий заказ на боевом бэкенде. Наличные — довести до доставки или отменить; карта — обязательно проверить возврат."),
])

# ============================ 2. Тест-план ============================
tp = wb.create_sheet("Тест-план")
tp.column_dimensions["A"].width = 34
tp.column_dimensions["B"].width = 115
tp["A1"] = "Тест-план: флоу доставки и отслеживания заказа (Customer Delivery Tracking)"
tp["A1"].font = TITLE
tp["A2"] = "Составлен 2026-09-30 по макету Figma «[RFD] Client • Mobile» и прогону 29.09 на сборке 1.1.8 (43)"
tp["A2"].font = NOTE
key_value_sheet(tp, 4, [
    ("1. Объект и цель", None),
    ("Что тестируем", "Экран отслеживания заказа в клиентском приложении (13 экранов макета: основной путь 01–06 и исключения 01A, 02A, 02B, 02C, 04A, 05A, 05B) и то, как тот же заказ видят мерчант, курьер и оператор."),
    ("Цель", "Каждый этап и каждое исключение корректно отображаются у клиента, данные совпадают на всех сторонах, деньги считаются и возвращаются правильно."),
    ("Вне области", "Каталог, корзина, профиль (есть в основном тест-плане TezChakana_TestPlan.xlsx), нагрузка, iOS."),
    ("2. Участники и устройства", None),
    ("[Клиент]", "Клиентское приложение новой сборки, язык RU. Эмулятор Android 17 или телефон (для сети и push — телефон)."),
    ("[Мерчант]", "tez-merchant.chakana.uz под суперадмином, в отдельном браузере или профиле (иначе вытеснит сессию оператора)."),
    ("[Курьер]", "Курьерское приложение на телефоне, курьер на линии."),
    ("[Оператор]", "tez-merchant.chakana.uz под ролью оператора (Chrome), раздел «Заказы», фильтр Pepsi Market: /orders?branch_id=79."),
    ("Магазин и товары", "Pepsi Market, самые дешёвые позиции. Для замены (RPL) — заказ из 3 позиций."),
    ("3. Соответствие экранов макета и статусов", None),
    ("01 Ждём подтверждения магазина", "Заказ оформлен. Панель: «Подтверждён» (29.09 — сразу после оформления, статуса «Ожидает» не было). Q-D01."),
    ("02 Магазин готовит заказ", "Мерчант: «Готовится», затем «Готов к доставке» (клиент остаётся на 02)."),
    ("02A Ищем курьера", "Заказ готов, курьер не принял дольше порога."),
    ("02B Некоторые товары закончились", "Мерчант пометил товары «нет в наличии»; у клиента таймер ответа."),
    ("02C Заказ готовится дольше", "Сборка идёт дольше порога."),
    ("03 Курьер забирает заказ", "Курьер принял заказ. Панель: ожидается «Назначен курьер» (29.09 не появился — BUG-D08)."),
    ("04 Курьер в пути", "Курьер забрал заказ. Панель: «Передан курьеру» / «В пути»."),
    ("04A Курьер задерживается", "Курьер не прибыл дольше порога; новое время доставки."),
    ("05 Курьер приехал", "Курьер «На месте», у клиента таймер ожидания. В панели такого статуса нет (Q-D02)."),
    ("05A Курьер не может с вами связаться", "Клиент недоступен (кнопка курьера или истечение ожидания)."),
    ("05B Не удалось доставить заказ", "Таймаут после 05A или решение. Статус панели — уточнить (Q-D14)."),
    ("06 Заказ доставлен", "Курьер передал заказ. Панель: «Доставлен»."),
    ("01A Заказ отменён", "Мерчант отклонил заказ. Панель: «Отменён»."),
    ("4. Порядок прогона по заказам", None),
    ("Заказ 1 — наличные", "Основной путь: OFR-01, OFR-02, TRK-01…TRK-05, OPR-01, OPR-02 → DLV-01, RSL-01, TRK-08, TRK-09 → DLV-02, CUR-01 → DLV-03, DLV-04, OPR-05 → RSL-04, DLV-05 → RSL-05, DLV-06, DLV-07, RSL-02, CUR-03 → DLV-08 → DLV-09, DLV-10, DLV-11, OPR-03. Попутно: TRK-06, TRK-07, OPR-04, RSL-06, CNL-04, CUR-02."),
    ("Заказ 2 — наличные", "Отмена клиентом: CNL-01, CNL-02."),
    ("Заказ 3 — карта", "Отклонение: OFR-03, REJ-01, REJ-03."),
    ("Заказ 4 — наличные", "Отклонение без возврата: REJ-02."),
    ("Заказ 5 — наличные, 3 позиции", "Замена с ответом: RPL-01, RPL-06, RPL-02, RPL-03 или RPL-04, RPL-05, RPL-10, RPL-08."),
    ("Заказ 6 — карта, 3 позиции", "Замена без ответа: RPL-01, RPL-07 → затем REJ-05 (мерчант отменяет → полный возврат)."),
    ("Заказ 7 — наличные", "Задержки и недоступный клиент: DLY-01, DLY-02, DLY-03, DLY-04, DLY-05, DLY-07, DLY-06 → UNV-01, UNV-02, CUR-04 → UNV-03. Попутно OPR-06."),
    ("Заказ 8 — карта", "Недоставка: до 05 → UNV-01 → UNV-04, UNV-05, CUR-05 → UNV-06, UNV-07."),
    ("По желанию", "CNL-03 (отмена клиентом с картой), REJ-04 (автоотмена), RPL-09 (закончилось всё), OPR-07…OPR-11, RSL-03, RSL-07, RSL-08."),
    ("5. Деньги", None),
    ("Наличные", "Заказы 1, 2, 4, 5, 7 — деньги не списываются; доводим до доставки или отмены."),
    ("Карта", "Заказы 3, 6, 8 — только там, где ожидается возврат. Каждый такой заказ довести до возврата и проверить поступление в банке; сумму и статус записать в «Комментарий»."),
    ("6. Безопасность", None),
    ("Боевой бэкенд", "Работаем только со своими тестовыми заказами. Чужие заказы в панели не открывать и не менять, их данные никуда не переносить."),
    ("Панель", "Ручную смену статуса (OPR-08) — только на своём тестовом заказе. По ссылке Telegram поддержки в панели не переходить (BUG-D09)."),
    ("Не оставлять хвостов", "Прерванный сценарий — отменить заказ как мерчант; для карты проверить возврат."),
    ("7. Критерии", None),
    ("Флоу годен", "Все P0 — Pass; нет открытых Critical; суммы у клиента, в панели и у курьера совпадают; возвраты по карте приходят."),
    ("Флоу не годен", "Любой Fail в Smoke; статус не доходит до клиента; неверные суммы или возвраты; заказ «застревает» на этапе."),
    ("Приоритет кейса", "P0 — без этого заказ не доставить или теряются деньги. P1 — важная ветка флоу. P2 — редкий случай, внешний вид."),
    ("Серьёзность бага", "Critical — деньги, безопасность, флоу не проходит. Major — неверные данные или этап, есть обход. Minor — текст, формат, мелкая логика. Trivial — внешний вид."),
])

# ============================ 3. Тест-кейсы ============================
tc = wb.create_sheet("Тест-кейсы")
tc_headers = ["ID", "Модуль", "Название", "Предусловия", "Шаги", "Ожидаемый результат",
              "Приоритет", "Тип", "Набор", "Риск", "Автотест", "Примечание / баг"]
header(tc, 1, tc_headers, [9, 14, 34, 32, 50, 60, 9, 6, 11, 14, 10, 34])
for i, case in enumerate(CASES, 2):
    body_row(tc, i, case, center_cols=(1, 7, 8, 9, 10, 11))
last_tc = len(CASES) + 1
tc.freeze_panes = "C2"
tc.auto_filter.ref = f"A1:L{last_tc}"
suite_colors = {"Smoke": RED, "Sanity": ORANGE, "Regression": BLUE}
risk_colors = {"Безопасно": GREEN, "Меняет данные": ORANGE, "Реальный заказ": RED, "Необратимо": PURPLE}
for val, col in suite_colors.items():
    tc.conditional_formatting.add(f"I2:I{last_tc}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill(col)))
for val, col in risk_colors.items():
    tc.conditional_formatting.add(f"J2:J{last_tc}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill(col)))
tc.conditional_formatting.add(f"G2:G{last_tc}", CellIsRule(operator="equal", formula=['"P0"'], font=Font(name=FONT, bold=True, color="C00000")))

sc = 14  # колонка N — сводка
tc.column_dimensions[get_column_letter(sc)].width = 18
tc.column_dimensions[get_column_letter(sc + 1)].width = 9
for col, text in ((sc, "Сводка"), (sc + 1, "Кол-во")):
    c = tc.cell(row=1, column=col, value=text)
    c.font, c.fill = HDR_FONT, HDR_FILL
summary = [
    ("Всего кейсов", f"=COUNTA(A2:A{last_tc})"),
    ("Smoke", f'=COUNTIF(I2:I{last_tc},"Smoke")'),
    ("Sanity (вкл. Smoke)", f'=COUNTIF(I2:I{last_tc},"Smoke")+COUNTIF(I2:I{last_tc},"Sanity")'),
    ("Regression (все)", f"=COUNTA(A2:A{last_tc})"),
    ("P0", f'=COUNTIF(G2:G{last_tc},"P0")'),
    ("P1", f'=COUNTIF(G2:G{last_tc},"P1")'),
    ("P2", f'=COUNTIF(G2:G{last_tc},"P2")'),
    ("Негативные (−)", f'=COUNTIF(H2:H{last_tc},"−")'),
    ("Реальный заказ", f'=COUNTIF(J2:J{last_tc},"Реальный заказ")'),
]
for j, (label, f) in enumerate(summary, 2):
    a = tc.cell(row=j, column=sc, value=label)
    b = tc.cell(row=j, column=sc + 1, value=f)
    a.font, b.font = BODY, BOLD
    a.border = b.border = BORDER

# ============================ 4. Прогон ============================
rn = wb.create_sheet("Прогон")
for i, w in enumerate([9, 14, 40, 9, 11, 14, 11, 12, 50], 1):
    rn.column_dimensions[get_column_letter(i)].width = w
rn["A1"] = "Прогон флоу доставки (шаблон) — скопируйте лист на каждую сборку"
rn["A1"].font = TITLE
meta = [("Сборка (версия)", "1.1.8 (43)"), ("Дата", ""), ("Тестировщик", ""), ("Устройство клиента", ""),
        ("Вид прогона", ""), ("Номера заказов", "")]
for j, (k, v) in enumerate(meta, 3):
    rn.cell(row=j, column=1, value=k).font = BOLD
    rn.merge_cells(start_row=j, start_column=2, end_row=j, end_column=3)
    c = rn.cell(row=j, column=2, value=v)
    c.fill, c.font, c.border = INPUT_FILL, BODY, BORDER
dv_kind = DataValidation(type="list", formula1='"Smoke,Sanity,Регресс,Ретест"', allow_blank=True)
rn.add_data_validation(dv_kind)
dv_kind.add("B7")

FIRST = 12
LAST = FIRST + len(CASES) - 1
G = f"$G${FIRST}:$G${LAST}"
E = f"$E${FIRST}:$E${LAST}"
stats = [
    ("Всего кейсов", f"=COUNTA($A${FIRST}:$A${LAST})"),
    ("Pass", f'=COUNTIF({G},"Pass")'),
    ("Fail", f'=COUNTIF({G},"Fail")'),
    ("Blocked", f'=COUNTIF({G},"Blocked")'),
    ("Skip", f'=COUNTIF({G},"Skip")'),
    ("Не выполнено", "=F3-F4-F5-F6-F7"),
    ("% Pass (из выполненных)", '=IF(F4+F5+F6=0,"-",F4/(F4+F5+F6))'),
]
for j, (k, f) in enumerate(stats, 3):
    a = rn.cell(row=j, column=5, value=k)
    b = rn.cell(row=j, column=6, value=f)
    a.font, b.font = BODY, BOLD
    a.border = b.border = BORDER
rn["F9"].number_format = "0%"
rn.column_dimensions["E"].width = 22
verdict = [
    ("H3", "Smoke выполнено", "I3", f'=COUNTIFS({E},"Smoke",{G},"<>")&" из "&COUNTIF({E},"Smoke")'),
    ("H4", "Smoke Fail", "I4", f'=COUNTIFS({E},"Smoke",{G},"Fail")'),
    ("H5", "P0 Fail", "I5", f'=COUNTIFS($D${FIRST}:$D${LAST},"P0",{G},"Fail")'),
    ("H6", "Вердикт", "I6", f'=IF(I4>0,"НЕ ГОДЕН: упал Smoke",IF(I5>0,"НЕ ГОДЕН: упал P0",IF(COUNTIFS({E},"Smoke",{G},"")>0,"Smoke не завершён","Smoke пройден")))'),
]
for lref, label, vref, f in verdict:
    rn[lref] = label
    rn[lref].font = BOLD if label == "Вердикт" else BODY
    rn[vref] = f
    rn[vref].font, rn[vref].border = BOLD, BORDER

header(rn, FIRST - 1, ["ID", "Модуль", "Название", "Приоритет", "Набор", "Риск", "Статус", "Баг", "Комментарий (№ заказа, факт)"])
TCR = f"'Тест-кейсы'!$A$2:$A${last_tc}"


def lookup(col):
    return f"=IFERROR(INDEX('Тест-кейсы'!${col}$2:${col}${last_tc},MATCH($A{{r}},{TCR},0)),\"\")"


tmpl = {2: lookup("B"), 3: lookup("C"), 4: lookup("G"), 5: lookup("I"), 6: lookup("J")}
dv_status = DataValidation(type="list", formula1='"Pass,Fail,Blocked,Skip"', allow_blank=True)
rn.add_data_validation(dv_status)
for k, case in enumerate(CASES):
    r = FIRST + k
    rn.cell(row=r, column=1, value=case[0])
    for col, f in tmpl.items():
        rn.cell(row=r, column=col, value=f.format(r=r))
    for col in range(1, 10):
        c = rn.cell(row=r, column=col)
        c.font, c.border = BODY, BORDER
        c.alignment = CENTER if col in (1, 4, 5, 6, 7, 8) else WRAP
    for col in (7, 8, 9):
        rn.cell(row=r, column=col).fill = INPUT_FILL
dv_status.add(f"G{FIRST}:G{LAST}")
status_colors = {"Pass": GREEN, "Fail": RED, "Blocked": ORANGE, "Skip": GREY}
for val, col in status_colors.items():
    rn.conditional_formatting.add(f"G{FIRST}:G{LAST}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill(col)))
for val, col in suite_colors.items():
    rn.conditional_formatting.add(f"E{FIRST}:E{LAST}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill(col)))
rn.freeze_panes = f"D{FIRST}"
rn.auto_filter.ref = f"A{FIRST-1}:I{LAST}"
rn["A10"] = "Жёлтые ячейки — заполнять. Пример строки: DLV-03 | Fail | BUG-D08 | «TEZ00871: в панели сразу «В пути», курьера в карточке нет»."
rn["A10"].font = NOTE

# ============================ 5. Баги ============================
bg = wb.create_sheet("Баги")
header(bg, 1, ["ID", "Модуль", "Заголовок", "Шаги", "Факт", "Ожидание", "Серьёзность", "Статус",
               "Найден в", "Кейс", "Дата", "Комментарий / трекер"],
       [10, 12, 40, 42, 45, 36, 11, 15, 22, 9, 11, 34])
for i, b in enumerate(BUGS, 2):
    body_row(bg, i, list(b[:10]) + ["2026-09-29", b[10]], center_cols=(1, 7, 8, 10, 11))
last_bug = len(BUGS) + 1
bg.freeze_panes = "C2"
bg.auto_filter.ref = f"A1:L{last_bug + 50}"
dv_sev = DataValidation(type="list", formula1='"Critical,Major,Minor,Trivial"', allow_blank=True)
dv_bst = DataValidation(type="list", formula1='"Новый,Перепроверить,Проверить на телефоне,Передан,Исправлен,Закрыт,Переоткрыт,Не баг"', allow_blank=True)
bg.add_data_validation(dv_sev)
bg.add_data_validation(dv_bst)
dv_sev.add(f"G2:G{last_bug + 50}")
dv_bst.add(f"H2:H{last_bug + 50}")
for val, col in {"Critical": "FF9C9C", "Major": RED, "Minor": ORANGE, "Trivial": GREY}.items():
    bg.conditional_formatting.add(f"G2:G{last_bug + 50}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill(col)))
bg.conditional_formatting.add(f"H2:H{last_bug + 50}", CellIsRule(operator="equal", formula=['"Закрыт"'], fill=fill(GREEN)))
bg.conditional_formatting.add(f"H2:H{last_bug + 50}", CellIsRule(operator="equal", formula=['"Переоткрыт"'], fill=fill(RED)))

# ============================ 6. Вопросы ============================
qs = wb.create_sheet("Вопросы")
header(qs, 1, ["ID", "Вопрос", "Связано с", "Ответ", "Дата ответа"], [8, 90, 26, 50, 12])
for i, q in enumerate(QUESTIONS, 2):
    body_row(qs, i, list(q) + ["", ""], center_cols=(1,))
    for col in (4, 5):
        qs.cell(row=i, column=col).fill = INPUT_FILL

for sheet in wb.worksheets:
    sheet.sheet_view.zoomScale = 110
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT, "cases", len(CASES), "bugs", len(BUGS), "questions", len(QUESTIONS))
