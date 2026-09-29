# -*- coding: utf-8 -*-
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import FormulaRule, CellIsRule
from openpyxl.utils import get_column_letter

sys.path.insert(0, sys.argv[1])
from cases import CASES, BUGS, QUESTIONS

OUT = sys.argv[2]
FONT = "Arial"
HDR_FILL = PatternFill("solid", fgColor="3A3A3A")
HDR_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
BOLD = Font(name=FONT, size=10, bold=True)
TITLE = Font(name=FONT, size=14, bold=True)
SUB = Font(name=FONT, size=11, bold=True)
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
SECTION_FILL = PatternFill("solid", fgColor="F2F2F2")
thin = Side(style="thin", color="D0D0D0")
BORDER = Border(left=thin, right=thin, top=thin, bottom=thin)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="top", wrap_text=True)

def fill(hex_):
    return PatternFill("solid", fgColor=hex_)

GREEN, RED, ORANGE, GREY, BLUE, PURPLE = "C6EFCE", "FFC7CE", "FFE0B2", "E7E6E6", "DDEBF7", "E4DFEC"

def header(ws, row, titles, widths=None):
    for i, t in enumerate(titles, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.alignment, c.border = HDR_FONT, HDR_FILL, CENTER, BORDER
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w

def body_row(ws, row, values, center_cols=()):
    for i, v in enumerate(values, 1):
        c = ws.cell(row=row, column=i, value=v)
        c.font, c.border = BODY, BORDER
        c.alignment = CENTER if i in center_cols else WRAP

wb = Workbook()

# ============================ 1. Как пользоваться ============================
ws = wb.active
ws.title = "Как пользоваться"
ws.column_dimensions["A"].width = 26
ws.column_dimensions["B"].width = 110
ws["A1"] = "Tez Chakana — ручное тестирование: как пользоваться этим файлом"
ws["A1"].font = TITLE
rows = [
    ("Листы", None),
    ("Тест-план", "Что тестируем, какие бывают прогоны (Smoke / Sanity / Регресс / Ретест / Исследовательское), когда какой запускать, критерии «сборка годна / не годна», правила безопасности для реального аккаунта."),
    ("Тест-кейсы", "Все кейсы. Фильтруйте по колонкам «Набор», «Модуль», «Приоритет», «Риск». Это «база» — правьте её, когда меняется продукт."),
    ("Прогон", "Шаблон одного прогона. На каждую сборку: ПКМ по вкладке → «Переместить или скопировать» → «Создать копию», переименовать (например «Прогон 1.1.9 Smoke»). Названия кейсов подтягиваются из листа «Тест-кейсы» по ID."),
    ("Баги", "Найденные баги с шагами. Статус меняйте по мере исправлений; исправленный баг обязательно перепроверьте (ретест)."),
    ("Вопросы", "Открытые вопросы к продукту/разработке. Ответ может превратить «вопрос» в баг или в новый кейс."),
    ("Наборы (колонка «Набор»)", None),
    ("Smoke", "Минимальная проверка, что сборка вообще работает и главный сценарий «найти → положить в корзину → заказать» проходит. ~30–40 мин. Фильтр: Набор = Smoke."),
    ("Sanity", "Проверка затронутых изменениями модулей. Фильтр: Модуль = <изменённый модуль>, Набор = Smoke или Sanity. Плюс ретест исправленных багов."),
    ("Regression", "Полный прогон: все кейсы (Smoke + Sanity + Regression). Можно разбить на 2–3 дня по модулям."),
    ("Вложенность", "Smoke ⊂ Sanity ⊂ Regression: в колонке указан самый маленький набор, куда кейс входит. Регресс = все строки."),
    ("Статусы в «Прогон»", None),
    ("Pass", "Фактический результат совпал с ожидаемым."),
    ("Fail", "Не совпал → завести/найти баг в листе «Баги» и вписать его ID в колонку «Баг»."),
    ("Blocked", "Нельзя выполнить из-за другого бага или окружения (нет сети, нет карты и т.п.) — причину в «Комментарий»."),
    ("Skip", "Сознательно пропущен (не входит в текущий прогон, нет данных)."),
    ("Пусто", "Ещё не выполнен."),
    ("Пример заполнения строки", "ORDH-03 | … | Статус: Fail | Баг: BUG-001 | Комментарий: «TEZ00812: 3 990 + 15 000, Jami 3 990»"),
    ("Что заполнять", "Жёлтые ячейки в листе «Прогон» (сборка, дата, устройство, вид прогона) и колонки «Статус», «Баг», «Комментарий». Остальное считается формулами."),
    ("Риск (колонка)", None),
    ("Безопасно", "Только просмотр, ничего не меняет."),
    ("Меняет данные", "Меняет корзину/адрес/язык/избранное и т.п. — после проверки вернуть как было."),
    ("Реальный заказ", "Создаёт настоящий заказ и может списать деньги. Сразу после проверки — отменить заказ в приложении."),
    ("Необратимо", "Удаление аккаунта и т.п. — только на отдельном тестовом аккаунте."),
    ("Колонка «Автотест»", "ID из docs/exploration-notes.md, если сценарий покрыт автотестом. После редизайна v1.1.8 автотесты требуют починки — пока на них не полагаться."),
]
r = 3
for k, v in rows:
    a = ws.cell(row=r, column=1, value=k)
    if v is None:
        a.font = SUB
        a.fill = SECTION_FILL
        ws.cell(row=r, column=2).fill = SECTION_FILL
    else:
        a.font = BOLD
        b = ws.cell(row=r, column=2, value=v)
        b.font, b.alignment = BODY, WRAP
    a.alignment = WRAP
    r += 1

# ============================ 2. Тест-план ============================
tp = wb.create_sheet("Тест-план")
tp.column_dimensions["A"].width = 30
tp.column_dimensions["B"].width = 115
tp["A1"] = "Тест-план: Tez Chakana (uz.agrobank.chakanaexpress), Android"
tp["A1"].font = TITLE
tp["A2"] = "Составлен 2026-09-29 по живому исследованию v1.1.8 (43) и заметкам docs/exploration-notes.md"
tp["A2"].font = Font(name=FONT, size=9, italic=True, color="666666")
plan = [
    ("1. Объект и цель", None),
    ("Продукт", "Мобильное приложение доставки из магазинов (продукты, цветы, кафе, кондитерские). Flutter, Android. Язык интерфейса по умолчанию — узбекский (латиница), есть Ўзбекча и Русский."),
    ("Цель", "Для каждой новой сборки быстро решить «годна / не годна», систематически находить регрессии в основных сценариях и вести баги в одном месте."),
    ("Модули", "Запуск, Вход, Home, Магазин, Поиск, Товар, Избранное, Корзина, Оформление, Заказ (реальные деньги), Заказы, Уведомления, Профиль, Мои данные, Адреса, Карты, Настройки, Общее (сеть, форматы, стабильность)."),
    ("Вне области", "Бэкенд/админка магазина и курьера, нагрузочное тестирование, безопасность, iOS (если появится — отдельный прогон по тем же кейсам)."),
    ("2. Окружение", None),
    ("Устройства", "Эмулятор Android 17 (1080×2400, sdk_gphone16k_arm64) — основной. Реальный Redmi M2006C3LG, Android 10, 720×1600 — Smoke на каждой сборке и всё, что связано с картой, push-уведомлениями, камерой, SMS."),
    ("Бэкенд", "Боевой. Аккаунт — реальный личный аккаунт владельца: реальные заказы, адреса, карты. См. раздел 6."),
    ("Вход", "Статического OTP больше нет — код из SMS вводится вручную. Чтобы не входить лишний раз, тесты с выходом (AUTH-08, ONB-05/06) выполняйте в конце прогона."),
    ("Эмулятор", "При нехватке памяти на Mac эмулятор даёт ложные зависания (ANR). Перед регрессом закройте тяжёлые приложения; при странных зависаниях — перезапуск эмулятора и повтор на телефоне."),
    ("3. Виды прогонов и когда их запускать", None),
    ("Smoke (≈30–40 мин)", "КАЖДАЯ новая сборка, до всего остального. 15 кейсов: установка поверх, запуск, Home, магазин, корзина, оформление, заказ наличными + отмена, заказы, профиль, версия. Любой Fail в Smoke = сборка «не годна», дальше не тестировать, сообщить разработчикам."),
    ("Ретест", "Каждая сборка, где заявлены исправления: пройти шаги каждого бага со статусом «Исправлен» из листа «Баги». Прошёл → «Закрыт», нет → «Переоткрыт»."),
    ("Sanity (≈30–60 мин)", "Хотфикс или небольшое изменение: Smoke + кейсы Sanity в изменённых модулях (фильтр Модуль + Набор ∈ {Smoke, Sanity}) + ретест. Если в изменениях написано «оформление» — прогнать Sanity модулей Корзина, Оформление, Заказ, Заказы."),
    ("Регресс (≈5–6 ч)", "Перед публикацией в сторе, после крупного обновления или редизайна, и не реже раза в 2–3 недели. Все кейсы. Можно разбить: день 1 — Запуск…Корзина, день 2 — Оформление…Заказы, день 3 — остальное."),
    ("Исследовательское (30–45 мин)", "На каждую новую фичу или сильно изменённый экран: сессия по заданной теме («что может сломаться в новом чекауте?»), заметки → баги и новые кейсы в лист «Тест-кейсы»."),
    ("4. Порядок работы на каждую сборку", None),
    ("Шаг 1", "Получить сборку и список изменений (changelog). Если changelog нет — попросить, без него Sanity превращается в угадывание."),
    ("Шаг 2", "Скопировать лист «Прогон», заполнить сборку, дату, устройство, вид прогона."),
    ("Шаг 3", "Установить поверх старой версии (ONB-04), проверить версию в настройках (SET-01)."),
    ("Шаг 4", "Smoke. Fail → стоп, отчёт разработчикам."),
    ("Шаг 5", "Ретест исправленных багов."),
    ("Шаг 6", "Sanity изменённых модулей (или Регресс, если релиз в стор / крупное изменение)."),
    ("Шаг 7", "Новые баги → лист «Баги»; новые сценарии → лист «Тест-кейсы»."),
    ("Шаг 8", "Отчёт: цифры из шапки листа «Прогон» + список новых/переоткрытых багов + решение «годна / не годна»."),
    ("5. Критерии", None),
    ("Вход в тестирование", "Сборка устанавливается и запускается; известна версия; есть changelog; есть доступ к аккаунту (SMS-код)."),
    ("Сборка годна к релизу", "Smoke 100% Pass; нет открытых Critical и Major багов в модулях Корзина / Оформление / Заказ / Заказы / Вход; все P0-кейсы Pass; ретест всех заявленных исправлений пройден."),
    ("Сборка не годна", "Любой Fail в Smoke; вылет или зависание на основном сценарии; неправильные суммы или списания; невозможно войти."),
    ("Приоритет кейса", "P0 — без этого нельзя заказать или теряются деньги. P1 — важная функция, есть обходной путь. P2 — второстепенное, внешний вид, редкие случаи."),
    ("Серьёзность бага", "Critical — деньги/суммы, вылет, нельзя оформить заказ или войти. Major — функция не работает или некорректные данные, есть обход. Minor — неудобство, неверный текст, мелкая логика. Trivial — внешний вид, форматирование."),
    ("6. Безопасность реального аккаунта", None),
    ("Заказы", "Реальные заказы разрешены владельцем при условии немедленной отмены в приложении (для карты — проверить возврат). Проверить актуальность списка магазинов, где оформлять нельзя (Palma cafe, Nur Rayhon, Mandarin, Tong Toragorgon, THE FEEL — от 2026-09-01), см. вопрос Q-05."),
    ("Данные", "Всё, что помечено «Меняет данные», вернуть как было: язык → O'zbekcha, адрес доставки, получатель, избранное, корзина пуста, тестовые адреса удалить."),
    ("Запрещено на основном аккаунте", "Удаление аккаунта (DET-08), удаление единственной рабочей карты, подтверждение отмены чужого/реального заказа без необходимости."),
    ("7. Шаблон баг-репорта", None),
    ("Заголовок", "Где + что не так, одной строкой: «Детали заказа: Jami не включает доставку»."),
    ("Поля", "Шаги (нумерованные, с данными) · Факт · Ожидание · Серьёзность · Версия/сборка · Устройство · Скриншот/видео · Связанный кейс."),
    ("Проверка перед отправкой", "Воспроизводится ли повторно? На телефоне тоже или только на эмуляторе? Может ли это увидеть обычный пользователь пальцем (а не только автотест)?"),
]
r = 4
for k, v in plan:
    a = tp.cell(row=r, column=1, value=k)
    if v is None:
        a.font = SUB
        for col in (1, 2):
            tp.cell(row=r, column=col).fill = SECTION_FILL
    else:
        a.font = BOLD
        b = tp.cell(row=r, column=2, value=v)
        b.font, b.alignment = BODY, WRAP
    a.alignment = WRAP
    r += 1

# ============================ 3. Тест-кейсы ============================
tc = wb.create_sheet("Тест-кейсы")
tc_headers = ["ID", "Модуль", "Название", "Предусловия", "Шаги", "Ожидаемый результат",
              "Приоритет", "Тип", "Набор", "Риск", "Автотест", "Примечание / баг"]
header(tc, 1, tc_headers, [10, 13, 34, 30, 48, 55, 9, 6, 11, 14, 11, 30])
for i, case in enumerate(CASES, 2):
    body_row(tc, i, case, center_cols=(1, 7, 8, 9, 10, 11))
last_tc = len(CASES) + 1
tc.freeze_panes = "C2"
tc.auto_filter.ref = f"A1:L{last_tc}"
suite_colors = {"Smoke": RED, "Sanity": ORANGE, "Regression": BLUE}
risk_colors = {"Безопасно": GREEN, "Меняет данные": ORANGE, "Реальный заказ": RED, "Необратимо": PURPLE}
for val, col in suite_colors.items():
    tc.conditional_formatting.add(f"I2:I{last_tc}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill(col)))
for val, col in risk_colors.items():
    tc.conditional_formatting.add(f"J2:J{last_tc}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill(col)))
tc.conditional_formatting.add(f"G2:G{last_tc}", CellIsRule(operator="equal", formula=['"P0"'], font=Font(name=FONT, bold=True, color="C00000")))

# summary block to the right
sc = 14  # column N
tc.column_dimensions[get_column_letter(sc)].width = 16
tc.column_dimensions[get_column_letter(sc + 1)].width = 9
tc.cell(row=1, column=sc, value="Сводка").font = HDR_FONT
tc.cell(row=1, column=sc).fill = HDR_FILL
tc.cell(row=1, column=sc + 1, value="Кол-во").font = HDR_FONT
tc.cell(row=1, column=sc + 1).fill = HDR_FILL
summary = [
    ("Всего кейсов", f"=COUNTA(A2:A{last_tc})"),
    ("Smoke", f'=COUNTIF(I2:I{last_tc},"Smoke")'),
    ("Sanity (вкл. Smoke)", f'=COUNTIF(I2:I{last_tc},"Smoke")+COUNTIF(I2:I{last_tc},"Sanity")'),
    ("Regression (все)", f"=COUNTA(A2:A{last_tc})"),
    ("P0", f'=COUNTIF(G2:G{last_tc},"P0")'),
    ("P1", f'=COUNTIF(G2:G{last_tc},"P1")'),
    ("P2", f'=COUNTIF(G2:G{last_tc},"P2")'),
    ("Негативные (−)", f'=COUNTIF(H2:H{last_tc},"−")'),
    ("Реальный заказ", f'=COUNTIF(J2:J{last_tc},"Реальный заказ")'),
]
for j, (label, f) in enumerate(summary, 2):
    a = tc.cell(row=j, column=sc, value=label)
    b = tc.cell(row=j, column=sc + 1, value=f)
    a.font, b.font = BODY, BOLD
    a.border = b.border = BORDER

# ============================ 4. Прогон ============================
rn = wb.create_sheet("Прогон")
widths = [10, 13, 40, 9, 11, 14, 11, 18, 50]
for i, w in enumerate(widths, 1):
    rn.column_dimensions[get_column_letter(i)].width = w
rn["A1"] = "Прогон (шаблон) — скопируйте лист на каждую сборку"
rn["A1"].font = TITLE
meta = [("Сборка (версия)", "1.1.8 (43)"), ("Дата", ""), ("Тестировщик", ""), ("Устройство", ""), ("Вид прогона", "")]
for j, (k, v) in enumerate(meta, 3):
    rn.cell(row=j, column=1, value=k).font = BOLD
    rn.merge_cells(start_row=j, start_column=2, end_row=j, end_column=3)
    c = rn.cell(row=j, column=2, value=v)
    c.fill, c.font, c.border = INPUT_FILL, BODY, BORDER
dv_kind = DataValidation(type="list", formula1='"Smoke,Sanity,Регресс,Ретест"', allow_blank=True)
rn.add_data_validation(dv_kind)
dv_kind.add("B7")

FIRST = 12
LAST = FIRST + len(CASES) - 1
G = f"$G${FIRST}:$G${LAST}"
E = f"$E${FIRST}:$E${LAST}"
stats = [
    ("Всего кейсов", f"=COUNTA($A${FIRST}:$A${LAST})"),
    ("Pass", f'=COUNTIF({G},"Pass")'),
    ("Fail", f'=COUNTIF({G},"Fail")'),
    ("Blocked", f'=COUNTIF({G},"Blocked")'),
    ("Skip", f'=COUNTIF({G},"Skip")'),
    ("Не выполнено", "=F3-F4-F5-F6-F7"),
    ("% Pass (из выполненных)", '=IF(F4+F5+F6=0,"-",F4/(F4+F5+F6))'),
]
for j, (k, f) in enumerate(stats, 3):
    a = rn.cell(row=j, column=5, value=k)
    b = rn.cell(row=j, column=6, value=f)
    a.font, b.font = BODY, BOLD
    a.border = b.border = BORDER
rn["F9"].number_format = "0%"
rn.column_dimensions["E"].width = 22
# Smoke progress
rn["H3"] = "Smoke выполнено"
rn["H3"].font = BODY
rn["I3"] = f'=COUNTIFS({E},"Smoke",{G},"<>")&" из "&COUNTIF({E},"Smoke")'
rn["H4"] = "Smoke Fail"
rn["H4"].font = BODY
rn["I4"] = f'=COUNTIFS({E},"Smoke",{G},"Fail")'
rn["H5"] = "P0 Fail"
rn["H5"].font = BODY
rn["I5"] = f'=COUNTIFS($D${FIRST}:$D${LAST},"P0",{G},"Fail")'
rn["H6"] = "Вердикт"
rn["H6"].font = BOLD
rn["I6"] = f'=IF(I4>0,"НЕ ГОДНА: упал Smoke",IF(I5>0,"НЕ ГОДНА: упал P0",IF(COUNTIFS({E},"Smoke",{G},"")>0,"Smoke не завершён","Smoke пройден")))'
for ref in ("I3", "I4", "I5", "I6"):
    rn[ref].font = BOLD
    rn[ref].border = BORDER

run_headers = ["ID", "Модуль", "Название", "Приоритет", "Набор", "Риск", "Статус", "Баг", "Комментарий"]
header(rn, FIRST - 1, run_headers)
TCR = f"'Тест-кейсы'!$A$2:$A${last_tc}"
def lookup(col):
    return f"=IFERROR(INDEX('Тест-кейсы'!${col}$2:${col}${last_tc},MATCH($A{{r}},{TCR},0)),\"\")"
tmpl = {2: lookup("B"), 3: lookup("C"), 4: lookup("G"), 5: lookup("I"), 6: lookup("J")}
dv_status = DataValidation(type="list", formula1='"Pass,Fail,Blocked,Skip"', allow_blank=True)
rn.add_data_validation(dv_status)
for k, case in enumerate(CASES):
    r = FIRST + k
    rn.cell(row=r, column=1, value=case[0])
    for col, f in tmpl.items():
        rn.cell(row=r, column=col, value=f.format(r=r))
    for col in range(1, 10):
        c = rn.cell(row=r, column=col)
        c.font, c.border = BODY, BORDER
        c.alignment = CENTER if col in (1, 4, 5, 6, 7, 8) else WRAP
    for col in (7, 8, 9):
        rn.cell(row=r, column=col).fill = INPUT_FILL
dv_status.add(f"G{FIRST}:G{LAST}")
status_colors = {"Pass": GREEN, "Fail": RED, "Blocked": ORANGE, "Skip": GREY}
for val, col in status_colors.items():
    rn.conditional_formatting.add(f"G{FIRST}:G{LAST}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill(col)))
for val, col in suite_colors.items():
    rn.conditional_formatting.add(f"E{FIRST}:E{LAST}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill(col)))
rn.freeze_panes = f"D{FIRST}"
rn.auto_filter.ref = f"A{FIRST-1}:I{LAST}"
rn["A9"] = "Жёлтые ячейки — заполнять. Фильтр по «Набор»: Smoke / Smoke+Sanity нужного модуля / всё для регресса."
rn["A9"].font = Font(name=FONT, size=9, italic=True, color="666666")

# ============================ 5. Баги ============================
bg = wb.create_sheet("Баги")
bug_headers = ["ID", "Модуль", "Заголовок", "Шаги", "Факт", "Ожидание", "Серьёзность", "Статус",
               "Найден в", "Кейс", "Дата", "Комментарий / трекер"]
header(bg, 1, bug_headers, [9, 12, 38, 42, 45, 36, 11, 15, 22, 9, 11, 26])
for i, b in enumerate(BUGS, 2):
    body_row(bg, i, list(b) + ["2026-09-29", ""], center_cols=(1, 7, 8, 10, 11))
last_bug = len(BUGS) + 1
bg.freeze_panes = "C2"
bg.auto_filter.ref = f"A1:L{last_bug + 50}"
dv_sev = DataValidation(type="list", formula1='"Critical,Major,Minor,Trivial"', allow_blank=True)
dv_bst = DataValidation(type="list", formula1='"Новый,Перепроверить,Проверить на телефоне,Передан,Исправлен,Закрыт,Переоткрыт,Не баг"', allow_blank=True)
bg.add_data_validation(dv_sev)
bg.add_data_validation(dv_bst)
dv_sev.add(f"G2:G{last_bug + 50}")
dv_bst.add(f"H2:H{last_bug + 50}")
sev_colors = {"Critical": "FF9C9C", "Major": RED, "Minor": ORANGE, "Trivial": GREY}
for val, col in sev_colors.items():
    bg.conditional_formatting.add(f"G2:G{last_bug + 50}", CellIsRule(operator="equal", formula=[f'"{val}"'], fill=fill(col)))
bg.conditional_formatting.add(f"H2:H{last_bug + 50}", CellIsRule(operator="equal", formula=['"Закрыт"'], fill=fill(GREEN)))
bg.conditional_formatting.add(f"H2:H{last_bug + 50}", CellIsRule(operator="equal", formula=['"Переоткрыт"'], fill=fill(RED)))

# ============================ 6. Вопросы ============================
qs = wb.create_sheet("Вопросы")
header(qs, 1, ["ID", "Вопрос", "Связано с", "Ответ", "Дата ответа"], [7, 90, 22, 50, 12])
for i, q in enumerate(QUESTIONS, 2):
    body_row(qs, i, list(q) + ["", ""], center_cols=(1,))
    for col in (4, 5):
        qs.cell(row=i, column=col).fill = INPUT_FILL

for sheet in wb.worksheets:
    sheet.sheet_view.zoomScale = 110
from openpyxl.workbook.properties import CalcProperties
wb.calculation = CalcProperties(fullCalcOnLoad=True)
wb.save(OUT)
print("saved", OUT, "cases", len(CASES), "bugs", len(BUGS))
